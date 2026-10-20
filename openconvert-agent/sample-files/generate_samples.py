"""Generate real input and output files for testing OpenConvert conversions."""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from docx import Document

from openconvert.engine import (
    compress_pdf,
    convert,
    merge_pdfs,
    split_pdf,
)


ROOT = Path(__file__).resolve().parent
IMAGE_DIR = ROOT / "image-inputs"
PDF_DIR = ROOT / "pdf-inputs"
DOCUMENT_DIR = ROOT / "document-inputs"
OUTPUT_DIR = ROOT / "conversion-outputs"

WRITER_FORMATS = {
    ".doc", ".docx", ".docm", ".dot", ".dotx", ".dotm",
    ".odt", ".ott", ".odm", ".fodt", ".sxw", ".stw",
    ".rtf", ".epub",
}
CALC_FORMATS = {
    ".ods", ".ots", ".fods", ".xls", ".xlsx", ".xlsm",
    ".xlt", ".xltx", ".xltm",
}
IMPRESS_FORMATS = {
    ".odp", ".otp", ".fodp", ".ppt", ".pptx", ".pptm", ".pps", ".ppsx",
}
DRAW_FORMATS = {".odg", ".otg", ".fodg"}
TEXT_FORMATS = {".txt", ".html", ".htm", ".xhtml", ".csv"}


def create_images() -> None:
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (800, 500), "#e7f1ff")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 800, 100), fill="#173b66")
    draw.ellipse((80, 155, 300, 375), fill="#ffb347", outline="#c65d00", width=8)
    draw.rounded_rectangle(
        (360, 155, 710, 375), radius=28, fill="#55b88a", outline="#176b49", width=8
    )
    draw.text((32, 35), "OpenConvert sample image", fill="white")
    draw.text((105, 420), "Use this file to test image conversions.", fill="#173b66")

    for extension, file_format in [
        (".png", "PNG"),
        (".jpg", "JPEG"),
        (".jpeg", "JPEG"),
        (".webp", "WEBP"),
        (".bmp", "BMP"),
        (".tif", "TIFF"),
        (".tiff", "TIFF"),
    ]:
        image.save(IMAGE_DIR / f"sample-image{extension}", format=file_format)
    image.save(IMAGE_DIR / "sample-image.gif", format="GIF")
    image.close()


def create_pdfs() -> None:
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    for filename, text in [
        ("merge-first.pdf", "First PDF for the merge test."),
        ("merge-second.pdf", "Second PDF for the merge test."),
    ]:
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), text)
        document.save(PDF_DIR / filename)
        document.close()

    document = pymupdf.open()
    for page_number in range(1, 4):
        page = document.new_page()
        page.insert_text((72, 72), f"OpenConvert sample PDF - page {page_number}")
        page.insert_text(
            (72, 110),
            "Selectable text for PDF to Word, Excel, PowerPoint, and TXT tests.",
        )
    document.save(PDF_DIR / "sample-document.pdf")
    document.close()


def create_office_sources() -> dict[str, Path]:
    DOCUMENT_DIR.mkdir(parents=True, exist_ok=True)
    paths = {}

    word = Document()
    word.add_heading("OpenConvert document sample", level=1)
    word.add_paragraph("This document tests local conversion to PDF.")
    paths["writer"] = DOCUMENT_DIR / "sample-document.docx"
    word.save(paths["writer"])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Conversion sample"
    sheet.append(["Item", "Quantity", "Price"])
    sheet.append(["Sample item", 3, 12.5])
    paths["calc"] = DOCUMENT_DIR / "sample-spreadsheet.xlsx"
    workbook.save(paths["calc"])

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "OpenConvert presentation sample"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    textbox.text_frame.text = "This slide tests local conversion to PDF."
    paths["impress"] = DOCUMENT_DIR / "sample-presentation.pptx"
    presentation.save(paths["impress"])

    (DOCUMENT_DIR / "sample-text.txt").write_text(
        "OpenConvert text sample\nThis file tests conversion to PDF.\n", encoding="utf-8"
    )
    (DOCUMENT_DIR / "sample-page.html").write_text(
        "<!doctype html><html><body><h1>OpenConvert HTML sample</h1>"
        "<p>This file tests conversion to PDF.</p></body></html>\n",
        encoding="utf-8",
    )
    (DOCUMENT_DIR / "sample-page.htm").write_bytes(
        (DOCUMENT_DIR / "sample-page.html").read_bytes()
    )
    (DOCUMENT_DIR / "sample-page.xhtml").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<html xmlns="http://www.w3.org/1999/xhtml"><body><h1>OpenConvert XHTML sample</h1>'
        '<p>This file tests conversion to PDF.</p></body></html>',
        encoding="utf-8",
    )
    (DOCUMENT_DIR / "sample-table.csv").write_text(
        "Item,Quantity,Price\nSample item,3,12.50\n", encoding="utf-8"
    )
    return paths


def create_odm_sample() -> None:
    import zipfile

    namespaces = (
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
    )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<office:document-content {namespaces}office:version="1.2">'
        '<office:body><office:text><text:p>OpenConvert master document sample</text:p>'
        '</office:text></office:body></office:document-content>'
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<office:document-styles {namespaces}office:version="1.2">'
        '<office:styles/></office:document-styles>'
    )
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest '
        'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
        'manifest:version="1.2">'
        '<manifest:file-entry manifest:full-path="/" '
        'manifest:media-type="application/vnd.oasis.opendocument.text-master"/>'
        '<manifest:file-entry manifest:full-path="content.xml" '
        'manifest:media-type="text/xml"/>'
        '<manifest:file-entry manifest:full-path="styles.xml" '
        'manifest:media-type="text/xml"/>'
        '</manifest:manifest>'
    )
    path = DOCUMENT_DIR / "sample-document.odm"
    with zipfile.ZipFile(path, "w") as package:
        package.writestr(
            "mimetype",
            "application/vnd.oasis.opendocument.text-master",
            compress_type=zipfile.ZIP_STORED,
        )
        package.writestr("content.xml", content)
        package.writestr("styles.xml", styles)
        package.writestr("META-INF/manifest.xml", manifest)


def export_document_formats(sources: dict[str, Path]) -> None:
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError(
            "LibreOffice is required to generate all sample document formats. "
            "Install it and ensure `soffice` is available on PATH."
        )
    batches = [
        (sources["writer"], WRITER_FORMATS - {".odm"}),
        (sources["calc"], CALC_FORMATS),
        (sources["impress"], IMPRESS_FORMATS),
        (sources["impress"], DRAW_FORMATS),
    ]
    for source, extensions in batches:
        for extension in sorted(extensions):
            with tempfile.TemporaryDirectory(prefix="openconvert-sample-") as temporary:
                temp = Path(temporary)
                output_dir = temp / "output"
                output_dir.mkdir()
                profile = (temp / "profile").as_uri()
                result = subprocess.run(
                    [
                        soffice,
                        "--headless",
                        f"-env:UserInstallation={profile}",
                        "--convert-to",
                        extension.lstrip("."),
                        "--outdir",
                        str(output_dir),
                        str(source),
                    ],
                    capture_output=True,
                    text=True,
                    timeout=120,
                    check=False,
                )
                generated = output_dir / f"{source.stem}{extension}"
                if result.returncode or not generated.is_file():
                    detail = (result.stderr or result.stdout or "No output was created").strip()
                    raise RuntimeError(f"Cannot generate {extension} sample: {detail}")
                shutil.copy2(generated, DOCUMENT_DIR / generated.name)
    create_odm_sample()


def create_conversion_outputs() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    pdf = PDF_DIR / "sample-document.pdf"
    for target in ("png", "jpg", "jpeg", "docx", "xlsx", "pptx", "txt"):
        result = convert(pdf, target, OUTPUT_DIR)
        outputs = result if isinstance(result, list) else [result]
        print(f"{pdf.name} -> {', '.join(path.name for path in outputs)}")

    merged = merge_pdfs(
        [PDF_DIR / "merge-first.pdf", PDF_DIR / "merge-second.pdf"], OUTPUT_DIR
    )
    split = split_pdf(pdf, OUTPUT_DIR)
    compressed = compress_pdf(pdf, OUTPUT_DIR)
    print(f"merge -> {merged.name}")
    print(f"split -> {len(split)} one-page PDFs")
    print(f"compress -> {compressed.name}")


def main() -> None:
    create_images()
    create_pdfs()
    sources = create_office_sources()
    export_document_formats(sources)
    create_conversion_outputs()
    print(f"Sample inputs and outputs created under {ROOT}")


if __name__ == "__main__":
    main()
