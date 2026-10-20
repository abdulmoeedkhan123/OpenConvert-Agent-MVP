from pathlib import Path

import pymupdf
import pytest
from PIL import Image

from openconvert.engine import (
    ConversionError,
    compress_pdf,
    convert,
    convert_images_to_pdf,
    convert_pdf_to_office,
    convert_pdf_to_images,
    merge_pdfs,
    split_pdf,
)
from openconvert.planner import deterministic_plan


def make_pdf(path: Path, *texts: str) -> Path:
    document = pymupdf.open()
    for text in texts:
        page = document.new_page()
        page.insert_text((72, 72), text)
    document.save(path)
    document.close()
    return path


def test_selected_images_combine_into_pdf_without_overwriting(tmp_path: Path):
    first = tmp_path / "first.png"
    second = tmp_path / "second.png"
    Image.new("RGB", (20, 10), "red").save(first)
    Image.new("RGB", (10, 20), "blue").save(second)
    existing = tmp_path / "first.pdf"
    existing.write_bytes(b"keep me")

    output = convert_images_to_pdf([first, second])

    assert output.name == "first (1).pdf"
    assert existing.read_bytes() == b"keep me"
    with pymupdf.open(output) as document:
        assert document.page_count == 2


def test_pdf_to_images_exports_every_page(tmp_path: Path):
    source = make_pdf(tmp_path / "sample.pdf", "First page", "Second page")

    outputs = convert(source, "jpg")

    assert isinstance(outputs, list)
    assert [path.name for path in outputs] == [
        "sample-page-001.jpg",
        "sample-page-002.jpg",
    ]
    assert all(Image.open(path).format == "JPEG" for path in outputs)


def test_merge_split_and_compress_pdfs(tmp_path: Path):
    first = make_pdf(tmp_path / "first.pdf", "First")
    second = make_pdf(tmp_path / "second.pdf", "Second")

    merged = merge_pdfs([first, second])
    pages = split_pdf(merged)
    compressed = compress_pdf(merged)

    with pymupdf.open(merged) as document:
        assert document.page_count == 2
    assert len(pages) == 2
    assert all(pymupdf.open(page).page_count == 1 for page in pages)
    with pymupdf.open(compressed) as document:
        assert document.page_count == 2
    assert first.exists() and second.exists()


@pytest.mark.parametrize("target", ["docx", "xlsx", "pptx", "txt"])
def test_pdf_text_exports_create_nonempty_files(tmp_path: Path, target: str):
    source = make_pdf(tmp_path / "report.pdf", "Extractable report text")

    output = convert_pdf_to_office(source, target)

    assert output.suffix == f".{target}"
    assert output.stat().st_size > 0
    if target == "txt":
        assert "Extractable report text" in output.read_text(encoding="utf-8")
    elif target == "docx":
        from docx import Document

        assert "Extractable report text" in "\n".join(
            paragraph.text for paragraph in Document(output).paragraphs
        )
    elif target == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(output, read_only=True)
        assert workbook["Page 1"]["A1"].value == "Extractable report text"
        workbook.close()
    else:
        from pptx import Presentation

        presentation = Presentation(output)
        assert "Extractable report text" in "\n".join(
            shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame
        )


def test_planner_recognizes_pdf_to_docx(tmp_path: Path):
    source = make_pdf(tmp_path / "report.pdf", "Extractable report text")

    plan = deterministic_plan("convert this pdf to docx", [source])

    assert plan["action"] == "convert"
    assert plan["target"] == "docx"
    assert plan["files"] == [str(source)]


def test_scanned_pdf_without_text_has_clear_error(tmp_path: Path):
    source = make_pdf(tmp_path / "scan.pdf", "")

    with pytest.raises(ConversionError, match="Scanned PDFs need OCR"):
        convert_pdf_to_office(source, "docx")


def test_pdf_render_rejects_unsupported_image_target(tmp_path: Path):
    source = make_pdf(tmp_path / "sample.pdf", "Text")

    with pytest.raises(ConversionError, match="Unsupported PDF image target"):
        convert_pdf_to_images(source, "webp")
