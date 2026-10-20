"""Safe, local conversion recipes. Never modifies source files or overwrites outputs."""
from __future__ import annotations

import shutil
import subprocess
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageOps, UnidentifiedImageError

import pymupdf

IMAGE_FORMATS = {
    "png": "PNG",
    "jpg": "JPEG",
    "jpeg": "JPEG",
    "webp": "WEBP",
    "bmp": "BMP",
    "tif": "TIFF",
    "tiff": "TIFF",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
PDF_EXTENSIONS = {".pdf"}
DOCUMENT_EXTENSIONS = {
    ".doc", ".docx", ".docm", ".dot", ".dotx", ".dotm",
    ".odt", ".ott", ".odm", ".fodt", ".sxw", ".stw",
    ".rtf", ".txt", ".html", ".htm", ".xhtml", ".csv",
    ".odp", ".otp", ".fodp", ".ppt", ".pptx", ".pptm", ".pps", ".ppsx",
    ".ods", ".ots", ".fods", ".xls", ".xlsx", ".xlsm", ".xlt", ".xltx",
    ".xltm", ".odg", ".otg", ".fodg", ".epub",
}
PDF_OFFICE_TARGETS = {"docx", "xlsx", "pptx", "txt"}

class ConversionError(RuntimeError):
    """Raised when a conversion recipe cannot safely complete."""

def unique_output_path(path: Path) -> Path:
    """Return path if free, otherwise append a numeric suffix without overwriting."""
    if not path.exists():
        return path
    for number in range(1, 10000):
        candidate = path.with_name(f"{path.stem} ({number}){path.suffix}")
        if not candidate.exists():
            return candidate
    raise ConversionError(f"Could not find a free output name for {path.name}")

def _image_target_format(target: str) -> str:
    normalized = target.lower().lstrip(".")
    if normalized not in IMAGE_FORMATS:
        raise ConversionError(f"Unsupported image target: {target}")
    return normalized

def convert_image(source: Path, target: str, output_dir: Path | None = None) -> Path:
    source = source.expanduser().resolve()
    if not source.is_file():
        raise ConversionError(f"Input file does not exist: {source}")
    target = _image_target_format(target)
    output_dir = (output_dir or source.parent).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    output = unique_output_path(output_dir / f"{source.stem}.{target}")

    try:
        with Image.open(source) as original:
            image = ImageOps.exif_transpose(original)
            # JPEG does not support alpha; composite onto white instead of silently failing.
            if IMAGE_FORMATS[target] == "JPEG":
                if image.mode in ("RGBA", "LA") or (image.mode == "P" and "transparency" in image.info):
                    rgba = image.convert("RGBA")
                    background = Image.new("RGB", rgba.size, (255, 255, 255))
                    background.paste(rgba, mask=rgba.getchannel("A"))
                    image = background
                else:
                    image = image.convert("RGB")
            elif IMAGE_FORMATS[target] in ("BMP", "TIFF") and image.mode not in ("RGB", "L", "RGBA"):
                image = image.convert("RGB")
            save_options = {}
            if IMAGE_FORMATS[target] == "WEBP":
                save_options = {"quality": 90, "method": 4}
            image.save(output, format=IMAGE_FORMATS[target], **save_options)
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        # Do not remove anything that wasn't created by this invocation.
        if output.exists():
            output.unlink()
        raise ConversionError(f"Image conversion failed: {exc}") from exc

    if not output.is_file() or output.stat().st_size == 0:
        raise ConversionError("Conversion reported success but output is missing or empty.")
    return output

def _open_pdf(source: Path) -> pymupdf.Document:
    try:
        return pymupdf.open(source)
    except (pymupdf.FileDataError, RuntimeError, ValueError) as exc:
        raise ConversionError(f"Could not read PDF: {exc}") from exc

def _validate_pdf(source: Path) -> Path:
    source = source.expanduser().resolve()
    if not source.is_file():
        raise ConversionError(f"Input file does not exist: {source}")
    if source.suffix.lower() not in PDF_EXTENSIONS:
        raise ConversionError(f"Unsupported PDF input type: {source.suffix}")
    return source

def convert_image_to_pdf(source: Path, output_dir: Path | None = None) -> Path:
    source = source.expanduser().resolve()
    if not source.is_file():
        raise ConversionError(f"Input file does not exist: {source}")
    if source.suffix.lower() not in IMAGE_EXTENSIONS:
        raise ConversionError(f"Unsupported image input type: {source.suffix}")
    output_dir = (output_dir or source.parent).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    output = unique_output_path(output_dir / f"{source.stem}.pdf")
    try:
        with Image.open(source) as original:
            image = ImageOps.exif_transpose(original).convert("RGB")
            image.save(output, "PDF")
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        if output.exists():
            output.unlink()
        raise ConversionError(f"Image-to-PDF conversion failed: {exc}") from exc
    return output

def convert_images_to_pdf(sources: Iterable[Path], output_dir: Path | None = None) -> Path:
    sources = [source.expanduser().resolve() for source in sources]
    if not sources:
        raise ConversionError("Select at least one image to create a PDF.")
    for source in sources:
        if not source.is_file():
            raise ConversionError(f"Input file does not exist: {source}")
        if source.suffix.lower() not in IMAGE_EXTENSIONS:
            raise ConversionError(f"Unsupported image input type: {source.suffix}")

    output_dir = (output_dir or sources[0].parent).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    output = unique_output_path(output_dir / f"{sources[0].stem}.pdf")
    images: list[Image.Image] = []
    try:
        for source in sources:
            with Image.open(source) as original:
                images.append(ImageOps.exif_transpose(original).convert("RGB"))
        images[0].save(output, "PDF", save_all=True, append_images=images[1:])
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        if output.exists():
            output.unlink()
        raise ConversionError(f"Images-to-PDF conversion failed: {exc}") from exc
    finally:
        for image in images:
            image.close()
    return output

def convert_pdf_to_images(
    source: Path, target: str, output_dir: Path | None = None
) -> list[Path]:
    source = _validate_pdf(source)
    normalized = target.lower().lstrip(".")
    if normalized not in {"png", "jpg", "jpeg"}:
        raise ConversionError(f"Unsupported PDF image target: {target}")
    output_dir = (output_dir or source.parent).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    try:
        with _open_pdf(source) as document:
            if document.page_count == 0:
                raise ConversionError("The PDF has no pages to convert.")
            for index, page in enumerate(document, 1):
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                suffix = "jpg" if normalized == "jpeg" else normalized
                output = unique_output_path(
                    output_dir / f"{source.stem}-page-{index:03d}.{suffix}"
                )
                pixmap.save(output)
                outputs.append(output)
    except ConversionError:
        for output in outputs:
            output.unlink(missing_ok=True)
        raise
    except (pymupdf.FileDataError, RuntimeError, OSError, ValueError) as exc:
        for output in outputs:
            output.unlink(missing_ok=True)
        raise ConversionError(f"PDF-to-image conversion failed: {exc}") from exc
    return outputs

def merge_pdfs(sources: Iterable[Path], output_dir: Path | None = None) -> Path:
    sources = [_validate_pdf(source) for source in sources]
    if len(sources) < 2:
        raise ConversionError("Select at least two PDF files to merge.")
    output_dir = (output_dir or sources[0].parent).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    output = unique_output_path(output_dir / f"{sources[0].stem}-merged.pdf")
    merged = pymupdf.open()
    try:
        for source in sources:
            with _open_pdf(source) as document:
                merged.insert_pdf(document)
        if merged.page_count == 0:
            raise ConversionError("The selected PDFs contain no pages.")
        merged.save(output, garbage=4, deflate=True)
    except ConversionError:
        output.unlink(missing_ok=True)
        raise
    except (pymupdf.FileDataError, RuntimeError, OSError, ValueError) as exc:
        output.unlink(missing_ok=True)
        raise ConversionError(f"PDF merge failed: {exc}") from exc
    finally:
        merged.close()
    return output

def split_pdf(source: Path, output_dir: Path | None = None) -> list[Path]:
    source = _validate_pdf(source)
    output_dir = (output_dir or source.parent).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    try:
        with _open_pdf(source) as document:
            if document.page_count == 0:
                raise ConversionError("The PDF has no pages to split.")
            for index in range(document.page_count):
                page_document = pymupdf.open()
                try:
                    page_document.insert_pdf(document, from_page=index, to_page=index)
                    output = unique_output_path(
                        output_dir / f"{source.stem}-page-{index + 1:03d}.pdf"
                    )
                    page_document.save(output, garbage=4, deflate=True)
                    outputs.append(output)
                finally:
                    page_document.close()
    except ConversionError:
        for output in outputs:
            output.unlink(missing_ok=True)
        raise
    except (pymupdf.FileDataError, RuntimeError, OSError, ValueError) as exc:
        for output in outputs:
            output.unlink(missing_ok=True)
        raise ConversionError(f"PDF split failed: {exc}") from exc
    return outputs

def compress_pdf(source: Path, output_dir: Path | None = None) -> Path:
    source = _validate_pdf(source)
    output_dir = (output_dir or source.parent).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    output = unique_output_path(output_dir / f"{source.stem}-compressed.pdf")
    try:
        with _open_pdf(source) as document:
            if document.page_count == 0:
                raise ConversionError("The PDF has no pages to compress.")
            document.save(
                output, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True
            )
    except ConversionError:
        output.unlink(missing_ok=True)
        raise
    except (pymupdf.FileDataError, RuntimeError, OSError, ValueError) as exc:
        output.unlink(missing_ok=True)
        raise ConversionError(f"PDF compression failed: {exc}") from exc
    return output

def convert_pdf_to_office(
    source: Path, target: str, output_dir: Path | None = None
) -> Path:
    source = _validate_pdf(source)
    target = target.lower().lstrip(".")
    if target not in PDF_OFFICE_TARGETS:
        raise ConversionError(f"Unsupported PDF output type: {target}")
    output_dir = (output_dir or source.parent).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    output = unique_output_path(output_dir / f"{source.stem}.{target}")
    try:
        with _open_pdf(source) as document:
            pages = [page.get_text().strip() for page in document]
        if not any(pages):
            raise ConversionError(
                "No selectable text was found. Scanned PDFs need OCR, which is not supported yet."
            )
        if target == "txt":
            output.write_text("\n\n".join(pages), encoding="utf-8")
        elif target == "docx":
            from docx import Document

            word = Document()
            for index, text in enumerate(pages):
                if index:
                    word.add_page_break()
                for line in text.splitlines():
                    word.add_paragraph(line)
            word.save(output)
        elif target == "xlsx":
            from openpyxl import Workbook

            workbook = Workbook()
            workbook.remove(workbook.active)
            for index, text in enumerate(pages, 1):
                worksheet = workbook.create_sheet(f"Page {index}"[:31])
                for line_number, line in enumerate(text.splitlines(), 1):
                    worksheet.cell(row=line_number, column=1, value=line)
            workbook.save(output)
        else:
            from pptx import Presentation
            from pptx.util import Inches, Pt

            presentation = Presentation()
            layout = presentation.slide_layouts[6]
            for index, text in enumerate(pages, 1):
                slide = presentation.slides.add_slide(layout)
                slide.shapes.add_textbox(
                    Inches(0.5), Inches(0.25), Inches(9), Inches(0.5)
                ).text_frame.text = f"Page {index}"
                textbox = slide.shapes.add_textbox(
                    Inches(0.5), Inches(0.9), Inches(9), Inches(6)
                )
                textbox.text_frame.text = text
                for paragraph in textbox.text_frame.paragraphs:
                    paragraph.font.size = Pt(12)
            presentation.save(output)
    except ConversionError:
        output.unlink(missing_ok=True)
        raise
    except (pymupdf.FileDataError, RuntimeError, OSError, ValueError, ImportError) as exc:
        output.unlink(missing_ok=True)
        raise ConversionError(f"PDF-to-{target.upper()} conversion failed: {exc}") from exc
    return output

def convert_to_pdf(source: Path, output_dir: Path | None = None, timeout: int = 120) -> Path:
    """Convert a document to PDF using a locally installed LibreOffice CLI."""
    source = source.expanduser().resolve()
    if not source.is_file():
        raise ConversionError(f"Input file does not exist: {source}")
    if source.suffix.lower() not in DOCUMENT_EXTENSIONS:
        raise ConversionError(f"Unsupported document input type: {source.suffix}")
    soffice = shutil.which("soffice")
    if not soffice:
        candidates = [
            Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
            Path("/opt/homebrew/bin/soffice"),
            Path("/usr/local/bin/soffice"),
        ]
        soffice = next((str(p) for p in candidates if p.exists()), None)
    if not soffice:
        raise ConversionError("LibreOffice was not found. Install it and ensure `soffice` is available.")

    output_dir = (output_dir or source.parent).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    expected = output_dir / f"{source.stem}.pdf"

    # Convert into a private temporary directory, then move into a non-overwriting final path.
    import tempfile
    with tempfile.TemporaryDirectory(prefix="openconvert-") as tmp:
        tmp_path = Path(tmp)
        cmd = [
            soffice,
            "--headless",
            "-env:UserInstallation=file://" + str(tmp_path / "lo-profile"),
            "--convert-to", "pdf",
            "--outdir", str(tmp_path),
            str(source),
        ]
        try:
            result = subprocess.run(
                cmd, capture_output=True, text=True, timeout=timeout, check=False
            )
        except subprocess.TimeoutExpired as exc:
            raise ConversionError(f"Document conversion timed out after {timeout} seconds.") from exc
        generated = tmp_path / f"{source.stem}.pdf"
        if result.returncode != 0 or not generated.is_file() or generated.stat().st_size == 0:
            detail = (result.stderr or result.stdout or "Unknown LibreOffice error").strip()
            raise ConversionError(f"Document conversion failed: {detail}")
        output = unique_output_path(expected)
        shutil.move(str(generated), str(output))

    return output

def supported_targets(source: Path) -> list[str]:
    """Return targets that this MVP can handle for a given source."""
    suffix = source.suffix.lower()
    if suffix in IMAGE_EXTENSIONS:
        return ["pdf", "png", "jpg", "jpeg", "webp", "bmp", "tiff"]
    if suffix in PDF_EXTENSIONS:
        return ["png", "jpg", "jpeg", *sorted(PDF_OFFICE_TARGETS)]
    if suffix in DOCUMENT_EXTENSIONS:
        return ["pdf"]
    return []

def convert(source: Path, target: str, output_dir: Path | None = None) -> Path | list[Path]:
    normalized = target.lower().lstrip(".")
    suffix = source.suffix.lower()
    if suffix in PDF_EXTENSIONS and normalized in {"png", "jpg", "jpeg"}:
        return convert_pdf_to_images(source, normalized, output_dir)
    if suffix in PDF_EXTENSIONS and normalized in PDF_OFFICE_TARGETS:
        return convert_pdf_to_office(source, normalized, output_dir)
    if suffix in IMAGE_EXTENSIONS and normalized == "pdf":
        return convert_image_to_pdf(source, output_dir)
    if normalized == "pdf":
        return convert_to_pdf(source, output_dir)
    return convert_image(source, normalized, output_dir)
